"""Workbook builder. Import from tools/build_hogan.py."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


def _styles():
    return dict(
        thin=Border(
            left=Side(style="thin", color="B0B0B0"),
            right=Side(style="thin", color="B0B0B0"),
            top=Side(style="thin", color="B0B0B0"),
            bottom=Side(style="thin", color="B0B0B0"),
        ),
        thick=Border(
            left=Side(style="medium", color="333333"),
            right=Side(style="medium", color="333333"),
            top=Side(style="medium", color="333333"),
            bottom=Side(style="medium", color="333333"),
        ),
        font=Font(name="Yu Gothic", size=10, color="222222"),
        font_b=Font(name="Yu Gothic", size=10, bold=True, color="222222"),
        font_title=Font(name="Yu Gothic", size=16, bold=True, color="111111"),
        font_sub=Font(name="Yu Gothic", size=11, bold=True, color="333333"),
        font_blue=Font(name="Yu Gothic", size=10, color="0000FF"),
        font_hint=Font(name="Yu Gothic", size=8, italic=True, color="666666"),
        font_head=Font(name="Yu Gothic", size=10, bold=True, color="FFFFFF"),
        fill_head=PatternFill("solid", fgColor="1F4E79"),
        fill_grid=PatternFill("solid", fgColor="F7F7F4"),
        fill_alt=PatternFill("solid", fgColor="EEF3F8"),
        fill_yel=PatternFill("solid", fgColor="FFF3B0"),
        fill_gov=PatternFill("solid", fgColor="F4C7C3"),
        fill_bnd=PatternFill("solid", fgColor="FCE4B3"),
        fill_nrm=PatternFill("solid", fgColor="D9EAD3"),
        fill_stb=PatternFill("solid", fgColor="D0E2F3"),
        fill_emg=PatternFill("solid", fgColor="FFF2CC"),
        center=Alignment(horizontal="center", vertical="center", wrap_text=True),
        left=Alignment(horizontal="left", vertical="center", wrap_text=True),
    )


def _a4_landscape(ws, header: str):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = header
    ws.oddFooter.left.text = "DebugShrine PUBLIC v1.0  |  hexagram weights = operator"
    ws.oddFooter.right.text = "Page &P / &N"


def build_workbook(path: str | Path) -> Path:
    S = _styles()
    wb = Workbook()
    _cover(wb.active, S)
    _flow(wb.create_sheet("01_FLOW"), S)
    _record(wb.create_sheet("02_RECORD"), S)
    _palace(wb.create_sheet("03_PALACE9"), S)
    _surface(wb.create_sheet("04_SURFACE"), S)
    _class(wb.create_sheet("05_CLASS"), S)
    _adapter(wb.create_sheet("06_ADAPTER"), S)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _cover(ws, S):
    ws.title = "00_COVER"
    ws.sheet_view.showGridLines = False
    _a4_landscape(ws, "00_COVER")
    for i in range(1, 21):
        ws.column_dimensions[get_column_letter(i)].width = 8
    ws.merge_cells("B2:P3")
    ws["B2"] = "DebugShrine PUBLIC v1.0  —  Excel方眼紙"
    ws["B2"].font = S["font_title"]
    ws.merge_cells("B5:P6")
    ws["B5"] = "Observe → Record → Aggregate → Compress(hook) → Classify"
    ws["B5"].font = S["font_sub"]
    blocks = [
        (8, "公開完成", "集計 + 分類 + 外部I/O + データフロー + 本ブック生成器"),
        (11, "非同梱", "局所コード生成 / Gate辞書 / 64卦の重み"),
        (14, "スナップショット", "既存ツールをカスタム。重み付けと運用は自己責任。"),
        (17, "正本", "Record。Sheetの面・卦は派生ビュー。"),
    ]
    for row, title, body in blocks:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=16)
        ws.cell(row, 2, title).font = S["font_b"]
        ws.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=16)
        ws.cell(row + 1, 2, body).font = S["font"]
    ws["B8"].fill = S["fill_nrm"]
    ws["B11"].fill = S["fill_yel"]
    ws.merge_cells("B21:P22")
    ws["B21"] = "01_FLOW / 02_RECORD / 03_PALACE9 / 04_SURFACE / 05_CLASS / 06_ADAPTER"
    ws["B21"].font = S["font_hint"]


def _flow(ws, S):
    _a4_landscape(ws, "01_FLOW")
    ws.sheet_view.showGridLines = True
    for i in range(1, 8):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.merge_cells("A1:F1")
    ws["A1"] = "外部連携データフロー"
    ws["A1"].font = S["font_title"]
    headers = ["段", "主体", "入力", "処理", "出力", "禁止"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(3, i, h)
        c.font = S["font_head"]
        c.fill = S["fill_head"]
        c.alignment = S["center"]
        c.border = S["thin"]
    rows = [
        ("0 Ingress", "CI/log/sheet/HTTP", "生イベント", "フィールド写経", "AuditRecord", "Classを書かない"),
        ("1 Record", "store", "AuditRecord", "Route* 導出", "facts", "物語ラベル"),
        ("2 Aggregate", "TrajectoryAggregator", "E,A", "node/edge++", "surface", "分類の先取り"),
        ("3 Compress", "external hook", "surface+counts", "既存卦ツール", "opaque Snapshot", "判定として使う"),
        ("4 Classify", "classify()", "Flag+RouteError", "oneOf 3値", "3 Classes", "局所コードを出す"),
        ("5 Egress", "JSONL/xlsx/git", "派生値", "ビュー出力", "写し", "正本化"),
    ]
    for r, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = S["font"]
            cell.alignment = S["center"]
            cell.border = S["thin"]
            cell.fill = S["fill_alt"] if r % 2 == 0 else S["fill_grid"]
        ws.row_dimensions[r].height = 28


def _record(ws, S):
    _a4_landscape(ws, "02_RECORD")
    ws.sheet_view.showGridLines = True
    ws.print_title_rows = "1:3"
    ws.auto_filter.ref = "A3:M22"
    ws.freeze_panes = "A4"
    titles = [
        "RecordId", "Timestamp", "Turn", "Intent",
        "ExpectedGate", "ActualGate", "GovernanceFlag",
        "RouteError", "RouteDelta", "Class", "Notes", "Valid",
        "EdgeFreq",
    ]
    widths = [14, 22, 8, 16, 16, 16, 16, 12, 16, 14, 18, 10, 12]
    ws.merge_cells("A1:M1")
    ws["A1"] = "Record 方眼紙 — 青字=入力。H–J,L は数式。"
    ws["A1"].font = S["font_title"]
    ws.merge_cells("A2:M2")
    ws["A2"] = "DATA_STRUCTURE.md 参照。行14–22は空スロット。"
    ws["A2"].font = S["font_hint"]
    for i, (h, w) in enumerate(zip(titles, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(3, i, h)
        c.font = S["font_head"]
        c.fill = S["fill_head"]
        c.alignment = S["center"]
        c.border = S["thin"]
    samples = [
        ("REC_001", "2026-09-30T00:00:00Z", "T1", "review", "A", "A", False),
        ("REC_002", "2026-09-30T00:01:00Z", "T2", "review", "A", "B", False),
        ("REC_003", "2026-09-30T00:02:00Z", "T3", "freeze", "A", "B", True),
        ("REC_004", "2026-09-30T00:03:00Z", "T4", "review", "A", "B", False),
        ("REC_005", "2026-09-30T00:04:00Z", "T5", "review", "C", "C", False),
        ("REC_006", "2026-09-30T00:05:00Z", "T6", "review", "A", "B", False),
        ("REC_007", "2026-09-30T00:06:00Z", "T7", "override", "B", "B", True),
        ("REC_008", "2026-09-30T00:07:00Z", "T8", "review", "A", "B", False),
        ("REC_009", "2026-09-30T00:08:00Z", "T9", "review", "A", "D", False),
        ("REC_010", "2026-09-30T00:09:00Z", "T10", "review", "A", "B", False),
    ]
    dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G4:G22")
    for i, s in enumerate(samples, 4):
        for col, val in enumerate(s, 1):
            cell = ws.cell(i, col, val)
            cell.font = S["font_blue"]
        _record_formulas(ws, i, S)
    for i in range(14, 23):
        for c in range(1, 8):
            cell = ws.cell(i, c, None)
            cell.font = S["font_blue"]
            cell.fill = S["fill_yel"]
            cell.border = S["thin"]
        _record_formulas(ws, i, S)
    ws.conditional_formatting.add("J4:J22", FormulaRule(formula=['J4="GOVERNANCE"'], fill=S["fill_gov"]))
    ws.conditional_formatting.add("J4:J22", FormulaRule(formula=['J4="BOUNDARY"'], fill=S["fill_bnd"]))
    ws.conditional_formatting.add("J4:J22", FormulaRule(formula=['J4="NORMAL"'], fill=S["fill_nrm"]))


def _record_formulas(ws, r, S):
    ws.cell(r, 8, f'=IF(OR(E{r}="",F{r}=""),"",E{r}<>F{r})')
    ws.cell(r, 9, f'=IF(OR(E{r}="",F{r}=""),"",E{r}&"→"&F{r})')
    ws.cell(
        r,
        10,
        f'=IF(G{r}="","",IF(G{r}=TRUE,"GOVERNANCE",IF(H{r}=TRUE,"BOUNDARY","NORMAL")))',
    )
    ws.cell(
        r,
        12,
        f'=IF(OR(E{r}="",F{r}="",G{r}=""),"WAIT",'
        f'IF(AND(J{r}="GOVERNANCE",G{r}=TRUE),"OK",'
        f'IF(AND(J{r}="BOUNDARY",G{r}=FALSE,H{r}=TRUE),"OK",'
        f'IF(AND(J{r}="NORMAL",G{r}=FALSE,H{r}=FALSE),"OK","FAIL"))))',
    )
    ws.cell(r, 13, f'=IF(I{r}="","",COUNTIF($I$4:$I$22,I{r}))')
    for c in range(1, 14):
        ws.cell(r, c).alignment = S["center"]
        ws.cell(r, c).border = S["thin"]
        if c >= 8:
            ws.cell(r, c).font = S["font"]


def _palace(ws, S):
    _a4_landscape(ws, "03_PALACE9")
    ws.sheet_view.showGridLines = True
    ws.merge_cells("A1:I1")
    ws["A1"] = "汎用 3×3 — セル意味は未定義。ActualGate が 1–9 のときだけカウント"
    ws["A1"].font = S["font_title"]
    for i in range(1, 10):
        ws.column_dimensions[get_column_letter(i)].width = 14
    n = 1
    for i in range(3):
        ws.row_dimensions[3 + i * 2].height = 18
        ws.row_dimensions[4 + i * 2].height = 36
        for j in range(3):
            lab = ws.cell(3 + i * 2, 2 + j * 2, f"cell {n}")
            lab.font = S["font_hint"]
            lab.alignment = S["center"]
            lab.border = S["thin"]
            cnt = ws.cell(4 + i * 2, 2 + j * 2, f'=COUNTIF(\'02_RECORD\'!F4:F22,"{n}")')
            cnt.font = S["font_b"]
            cnt.alignment = S["center"]
            cnt.border = S["thick"]
            cnt.fill = S["fill_stb"]
            n += 1
    ws.merge_cells("A12:I14")
    ws["A12"] = "サンプルゲートが A/B/C なら 0 のままが正しい。対応表は公開層に置かない。"
    ws["A12"].font = S["font"]
    ws["A12"].fill = S["fill_yel"]


def _surface(ws, S):
    _a4_landscape(ws, "04_SURFACE")
    ws.sheet_view.showGridLines = True
    ws.merge_cells("A1:H1")
    ws["A1"] = "軌道集計と面バンド"
    ws["A1"].font = S["font_title"]
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 40
    labels = [
        (3, "total", "=COUNTA('02_RECORD'!E4:E22)"),
        (4, "distinct_edges",
         '=SUMPRODUCT((\'02_RECORD\'!I4:I22<>"")/COUNTIF(\'02_RECORD\'!I4:I22,\'02_RECORD\'!I4:I22&""))'),
        (5, "STABLE_LT", 0.35),
        (6, "LOCKED_GE", 0.60),
    ]
    for r, name, val in labels:
        ws.cell(r, 1, name).font = S["font_b"]
        ws.cell(r, 1).border = S["thin"]
        cell = ws.cell(r, 2, val)
        cell.border = S["thin"]
        cell.alignment = S["center"]
    ws["B5"].font = S["font_blue"]
    ws["B5"].fill = S["fill_yel"]
    ws["B6"].font = S["font_blue"]
    ws["B6"].fill = S["fill_yel"]
    ws["C5"] = "公開デフォルト。変更は自己責任。"
    ws["C5"].font = S["font_hint"]
    for i, h in enumerate(["RouteDelta", "count", "ratio"], 1):
        c = ws.cell(8, i, h)
        c.font = S["font_head"]
        c.fill = S["fill_head"]
        c.alignment = S["center"]
        c.border = S["thin"]
    ws.merge_cells("A8:G9")
    ws["A8"] = "UNIQUE/FILTER は使わない。max_count は 02_RECORD!M4:M22 (EdgeFreq=COUNTIF) の最大。"
    ws["A8"].font = S["font_hint"]
    ws["A19"] = "max_count"
    ws["B19"] = "=IFERROR(MAX('02_RECORD'!M4:M22),0)"
    ws["A20"] = "dominant_ratio"
    ws["B20"] = "=IF(B3=0,0,B19/B3)"
    ws["B20"].number_format = "0.00"
    ws["A21"] = "Surface"
    ws["B21"] = '=IF(B3=0,"",IF(B20<B5,"Stable",IF(B20<B6,"Emerging","Locked")))'
    ws["B21"].font = S["font_b"]
    ws["B21"].fill = S["fill_emg"]
    for r in range(19, 22):
        ws.cell(r, 1).font = S["font_b"]
        ws.cell(r, 1).border = S["thin"]
        ws.cell(r, 2).border = S["thin"]
    ws["E3"] = "Class"
    ws["F3"] = "n"
    ws["E3"].font = S["font_head"]
    ws["F3"].font = S["font_head"]
    ws["E3"].fill = S["fill_head"]
    ws["F3"].fill = S["fill_head"]
    for i, name in enumerate(["GOVERNANCE", "BOUNDARY", "NORMAL"], 4):
        ws.cell(i, 5, name)
        ws.cell(i, 6, f'=COUNTIF(\'02_RECORD\'!J4:J22,"{name}")')
        for c in range(5, 7):
            ws.cell(i, c).border = S["thin"]
            ws.cell(i, c).alignment = S["center"]
    chart = BarChart()
    chart.type = "col"
    chart.title = "Class counts"
    chart.y_axis.title = "n"
    data = Reference(ws, min_col=6, min_row=3, max_row=6)
    cats = Reference(ws, min_col=5, min_row=4, max_row=6)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.y_axis.scaling.min = 0
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.style = 10
    chart.width = 12
    chart.height = 6
    ws.add_chart(chart, "E8")


def _class(ws, S):
    _a4_landscape(ws, "05_CLASS")
    ws.sheet_view.showGridLines = True
    ws.merge_cells("A1:G1")
    ws["A1"] = "分類表（公開・排他）"
    ws["A1"].font = S["font_title"]
    for i, h in enumerate(["GovernanceFlag", "RouteError", "Class", "意味"], 1):
        c = ws.cell(3, i, h)
        c.font = S["font_head"]
        c.fill = S["fill_head"]
        c.alignment = S["center"]
        c.border = S["thin"]
    triples = [
        (True, "any", "GOVERNANCE", "制御してよいか", S["fill_gov"]),
        (False, True, "BOUNDARY", "どこへ送ったか", S["fill_bnd"]),
        (False, False, "NORMAL", "残余", S["fill_nrm"]),
    ]
    for i, row in enumerate(triples, 4):
        *vals, fill = row
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i, c, v)
            cell.font = S["font"]
            cell.alignment = S["center"]
            cell.border = S["thin"]
            cell.fill = fill
    for col, w in enumerate([18, 14, 16, 28], 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.merge_cells("A8:G10")
    ws["A8"] = "local_code = consumer.bind(Class, catalog)  —  公開出力は3値で停止。"
    ws["A8"].fill = S["fill_yel"]
    ws["A8"].alignment = S["left"]


def _adapter(ws, S):
    _a4_landscape(ws, "06_ADAPTER")
    ws.sheet_view.showGridLines = True
    ws.merge_cells("A1:H1")
    ws["A1"] = "64卦スナップショット — 既存ツールカスタム（自己責任）"
    ws["A1"].font = S["font_title"]
    ws.merge_cells("A3:H5")
    ws["A3"] = (
        "エンジン非同梱。surface と counts だけ渡す。返却文字列を Snapshot に貼る。\n"
        "重み付け・運用・解釈はオペレータ。卦は監査判定ではない。"
    )
    ws["A3"].alignment = S["left"]
    ws["A3"].fill = S["fill_yel"]
    ws.column_dimensions["A"].width = 36
    rows = [
        (7, "Surface (from 04)", "='04_SURFACE'!B21"),
        (8, "dominant_ratio", "='04_SURFACE'!B20"),
        (9, "Snapshot (paste)", None),
        (10, "Tool name / version", None),
        (11, "Operator accepts weights", False),
    ]
    for r, label, val in rows:
        ws.cell(r, 1, label).font = S["font_b"]
        ws.cell(r, 1).border = S["thin"]
        cell = ws.cell(r, 2, val)
        cell.border = S["thin"]
        if val is None or val is False:
            cell.font = S["font_blue"]
            cell.fill = S["fill_yel"]
    ws.merge_cells("B9:H9")
    ws.merge_cells("B10:H10")
    ws["B8"].number_format = "0.00"
